"""Extract plain text from an uploaded resume, entirely on this computer.

The file is read in memory and never written to disk.
"""
from __future__ import annotations

import io

ALLOWED = {".pdf", ".docx", ".txt", ".md"}
MAX_CHARS = 20_000


class ResumeError(Exception):
    pass


def extract_text(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    ext = name[name.rfind("."):] if "." in name else ""
    if ext not in ALLOWED:
        raise ResumeError("Please upload a PDF, Word (.docx) or text file.")

    if ext == ".pdf":
        if not data.startswith(b"%PDF"):
            raise ResumeError("That file does not look like a PDF.")
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ResumeError("This PDF is password-protected.")
            text = "\n".join((page.extract_text() or "") for page in reader.pages[:10])
        except ResumeError:
            raise
        except Exception as e:
            raise ResumeError("Could not read this PDF.") from e
    elif ext == ".docx":
        if not data.startswith(b"PK"):
            raise ResumeError("That file does not look like a Word document.")
        from docx import Document

        try:
            doc = Document(io.BytesIO(data))
        except Exception as e:
            raise ResumeError("Could not read this Word document.") from e
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    else:
        text = data.decode("utf-8", errors="ignore")

    text = text.strip()
    if len(text) < 80:
        raise ResumeError(
            "Very little text was found. If this is a scanned PDF, please enter your details manually."
        )
    return text[:MAX_CHARS]
